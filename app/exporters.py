"""Сохранение расшифровки в TXT / SRT / VTT / DOCX."""

FORMATS = {
    "txt": "Текст (*.txt)",
    "srt": "Субтитры SRT (*.srt)",
    "vtt": "Субтитры WebVTT (*.vtt)",
    "docx": "Документ Word (*.docx)",
}


def _ts(seconds, sep=",", hours_always=True):
    ms = int(round(max(seconds, 0) * 1000))
    h, ms = divmod(ms, 3_600_000)
    m, ms = divmod(ms, 60_000)
    s, ms = divmod(ms, 1000)
    if hours_always or h:
        return f"{h:02d}:{m:02d}:{s:02d}{sep}{ms:03d}"
    return f"{m:02d}:{s:02d}{sep}{ms:03d}"


def short_ts(seconds):
    s = int(max(seconds, 0))
    h, s = divmod(s, 3600)
    m, s = divmod(s, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def paragraphs(segments, pause=1.5):
    """Склеивает сегменты в абзацы, разрывая на длинных паузах."""
    result, current, prev_end = [], [], None
    for seg in segments:
        text = seg["text"].strip()
        if not text:
            continue
        if current and prev_end is not None and seg["start"] - prev_end >= pause:
            result.append(" ".join(current))
            current = []
        current.append(text)
        prev_end = seg["end"]
    if current:
        result.append(" ".join(current))
    return result


def to_txt(item, timestamps=False):
    segs = item["segments"]
    if timestamps:
        return "\n".join(f"[{short_ts(s['start'])}] {s['text'].strip()}" for s in segs if s["text"].strip()) + "\n"
    return "\n\n".join(paragraphs(segs)) + "\n"


def to_srt(item):
    out = []
    for i, s in enumerate((s for s in item["segments"] if s["text"].strip()), 1):
        out.append(f"{i}\n{_ts(s['start'])} --> {_ts(s['end'])}\n{s['text'].strip()}\n")
    return "\n".join(out)


def to_vtt(item):
    out = ["WEBVTT\n"]
    for s in item["segments"]:
        if s["text"].strip():
            out.append(f"{_ts(s['start'], '.')} --> {_ts(s['end'], '.')}\n{s['text'].strip()}\n")
    return "\n".join(out)


def to_docx(item, path, timestamps=False):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.add_heading(item["name"], level=1)
    meta = doc.add_paragraph()
    run = meta.add_run(f"Длительность: {short_ts(item.get('duration') or 0)}   ·   "
                       f"Язык: {item.get('language', '—')}   ·   Модель: {item.get('model', '—')}")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    if timestamps:
        for s in item["segments"]:
            if not s["text"].strip():
                continue
            p = doc.add_paragraph()
            t = p.add_run(f"[{short_ts(s['start'])}]  ")
            t.font.color.rgb = RGBColor(0x7C, 0x5C, 0xFF)
            p.add_run(s["text"].strip())
    else:
        for para in paragraphs(item["segments"]):
            doc.add_paragraph(para)
    doc.save(path)


def export(item, fmt, path, timestamps=False):
    if fmt == "docx":
        to_docx(item, path, timestamps)
        return
    if fmt == "txt":
        text = to_txt(item, timestamps)
    elif fmt == "srt":
        text = to_srt(item)
    elif fmt == "vtt":
        text = to_vtt(item)
    else:
        raise ValueError(f"Неизвестный формат: {fmt}")
    # utf-8-sig, чтобы Блокнот и старые плееры корректно показывали кириллицу
    with open(path, "w", encoding="utf-8-sig", newline="\r\n" if fmt == "txt" else "\n") as f:
        f.write(text)
